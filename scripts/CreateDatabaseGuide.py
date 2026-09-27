from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from pathlib import Path

root=Path(r"C:\Users\Delta\Society360")
out=root/"Database"/"Society360_Database_Relationship_Guide.xlsx"

tables=[
("m_society","Master","Society master; root tenant","society_id","-"),
("m_society_branding","Master","Society branding","branding_id","society_id -> m_society"),
("m_user","Master","Login/user master","user_id","society_id -> m_society"),
("m_role","Master","Dynamic roles","role_id","-"),
("m_permission","Master","Dynamic module/action rights","permission_id","-"),
("m_role_permission","Security","Role to permission mapping","role_id, permission_id","role_id -> m_role; permission_id -> m_permission"),
("m_user_role","Security","User to role mapping","user_id, role_id","user_id -> m_user; role_id -> m_role"),
("m_user_society","Security","User society access mapping","user_id, society_id","user_id -> m_user; society_id -> m_society"),
("t_user_session","Security","Hashed session tokens","session_id","user_id -> m_user; society_id -> m_society"),
("m_module","Master","Application module catalog","module_id","-"),
("m_language","Master","Supported languages","language_code","-"),
("m_translation","Master","Multilingual resource text","translation_id","society_id -> m_society; language_code -> m_language"),
("m_system_setting","Configuration","Society/system settings","setting_id","society_id -> m_society"),
("m_building","Master","Buildings/towers","building_id","society_id -> m_society"),
("m_wing","Master","Wings","wing_id","society_id -> m_society; building_id -> m_building"),
("m_flat","Master","Flats/units","flat_id","society_id -> m_society; building_id -> m_building; wing_id -> m_wing"),
("m_customer","Master","Owners/tenants/residents","customer_id","society_id -> m_society"),
("m_customer_flat","Master","Customer-flat relationship","customer_flat_id","society_id -> m_society; customer_id -> m_customer; flat_id -> m_flat"),
("m_charge_type","Billing Master","Charge master","charge_type_id","society_id -> m_society"),
("m_rate_plan","Billing Master","Effective-dated rate plan","rate_plan_id","society_id -> m_society"),
("m_charge_rule","Billing Master","Charge calculation rules","charge_rule_id","society_id -> m_society; rate_plan_id -> m_rate_plan; charge_type_id -> m_charge_type"),
("m_interest_rule","Billing Master","DPC/interest rules","interest_rule_id","society_id -> m_society"),
("m_rebate_rule","Billing Master","Early payment rebate","rebate_rule_id","society_id -> m_society"),
("m_waiver_rule","Billing Master","Waiver rules","waiver_rule_id","society_id -> m_society"),
("m_charge_tax_rule","Billing Master","Tax/GST rules","tax_rule_id","society_id -> m_society"),
("m_payment_mode","Billing Master","Payment modes","payment_mode_id","society_id -> m_society"),
("m_parking_slot","Master","Parking slot master","parking_slot_id","society_id -> m_society"),
("m_vehicle","Master","Vehicle master","vehicle_id","society_id -> m_society; customer_id -> m_customer; flat_id -> m_flat; parking_slot_id -> m_parking_slot"),
("t_bill","Transaction","Billing header","bill_id","society_id -> m_society; flat_id -> m_flat"),
("t_bill_line_item","Transaction","Bill charge lines","bill_line_item_id","society_id -> m_society; bill_id -> t_bill; charge_type_id -> m_charge_type"),
("t_bill_status_history","Transaction","Bill lifecycle audit","history_id","society_id -> m_society; bill_id -> t_bill; changed_by -> m_user"),
("t_payment","Transaction","Collection/payment header","payment_id","society_id -> m_society; customer_id -> m_customer; flat_id -> m_flat"),
("t_payment_allocation","Transaction","Payment to bill allocation","allocation_id","society_id -> m_society; payment_id -> t_payment; bill_id -> t_bill"),
("t_receipt","Transaction","Receipt header","receipt_id","society_id -> m_society; payment_id -> t_payment"),
("t_collection_reversal","Transaction","Payment reversal","reversal_id","society_id -> m_society; payment_id -> t_payment"),
("t_adjustment","Transaction","Bill/flat adjustment","adjustment_id","society_id -> m_society; flat_id -> m_flat; bill_id -> t_bill; approved_by -> m_user"),
("t_complaint","Transaction","Complaint workflow","complaint_id","society_id -> m_society; flat_id -> m_flat; customer_id -> m_customer; assigned_to -> m_user"),
("t_visitor_entry","Transaction","Visitor/security entry","visitor_entry_id","society_id -> m_society; flat_id -> m_flat"),
("t_parking_assignment","Transaction","Parking assignment","assignment_id","society_id -> m_society; parking_slot_id -> m_parking_slot; flat_id -> m_flat; customer_id -> m_customer"),
("t_notice","Transaction","Society notices","notice_id","society_id -> m_society; created_by -> m_user"),
("t_notification","Transaction","SMS/email/app notifications","notification_id","society_id -> m_society; user_id -> m_user; customer_id -> m_customer"),
("t_document","Transaction","Resident/society documents","document_id","society_id -> m_society; flat_id -> m_flat; customer_id -> m_customer; uploaded_by -> m_user"),
("t_migration_batch","Transaction","Migration batch header","migration_batch_id","society_id -> m_society; created_by -> m_user"),
("t_migration_row","Transaction","Migration row tracking","migration_row_id","migration_batch_id -> t_migration_batch"),
("t_migration_error","Transaction","Migration validation errors","migration_error_id","migration_batch_id -> t_migration_batch"),
("t_audit_log","Audit","Business/security audit log","audit_log_id","society_id -> m_society; user_id -> m_user"),
]

relationships=[
("m_society","m_society_branding","1:1","society_id","Branding belongs to one society"),
("m_society","m_user","1:N","society_id","Society-scoped users; Super Admin may have NULL society"),
("m_society","m_building","1:N","society_id","Building isolation"),
("m_building","m_wing","1:N","building_id","Wing under building"),
("m_wing","m_flat","1:N","wing_id","Flat under wing"),
("m_society","m_flat","1:N","society_id","Tenant boundary"),
("m_customer","m_customer_flat","1:N","customer_id","Resident ownership/tenancy mapping"),
("m_flat","m_customer_flat","1:N","flat_id","Flat residents"),
("m_rate_plan","m_charge_rule","1:N","rate_plan_id","Effective charge rules"),
("m_charge_type","m_charge_rule","1:N","charge_type_id","Charge calculation"),
("m_flat","t_bill","1:N","flat_id","Bills generated for flats"),
("t_bill","t_bill_line_item","1:N","bill_id","Bill charges"),
("t_bill","t_payment_allocation","1:N","bill_id","Payments allocated to bills"),
("t_payment","t_payment_allocation","1:N","payment_id","Payment split across bills"),
("t_payment","t_receipt","1:1","payment_id","Receipt for payment"),
("m_flat","t_payment","1:N","flat_id","Collections against flat"),
("m_flat","t_complaint","1:N","flat_id","Flat complaints"),
("m_user","t_complaint","1:N","assigned_to","Complaint assignment"),
("m_flat","t_visitor_entry","1:N","flat_id","Visitors for flat"),
("m_parking_slot","m_vehicle","1:N","parking_slot_id","Vehicle parking"),
("m_customer","m_vehicle","1:N","customer_id","Vehicle owner"),
("m_user","m_user_role","1:N","user_id","User roles"),
("m_role","m_role_permission","1:N","role_id","Role rights"),
("m_permission","m_role_permission","1:N","permission_id","Permission assignment"),
("m_user","m_user_society","1:N","user_id","Society access"),
("m_society","m_user_society","1:N","society_id","User access"),
("m_user","t_user_session","1:N","user_id","Login sessions"),
]

modules=[
("Authentication","m_user, t_user_session, m_role, m_permission","Login, session, roles"),
("Multi-Society","m_society, m_user_society","Society isolation/access"),
("Buildings & Flats","m_building, m_wing, m_flat","Property master"),
("Residents","m_customer, m_customer_flat","Owner/tenant master"),
("Billing Configuration","m_charge_type, m_rate_plan, m_charge_rule, m_interest_rule, m_rebate_rule, m_waiver_rule, m_charge_tax_rule","Dynamic billing rules"),
("Billing","t_bill, t_bill_line_item, t_bill_status_history","Bill lifecycle"),
("Collection","t_payment, t_payment_allocation, t_receipt, t_collection_reversal, t_adjustment","Payment lifecycle"),
("Parking","m_parking_slot, m_vehicle, t_parking_assignment","Parking/vehicles"),
("Complaints","t_complaint","Complaint workflow"),
("Security","t_visitor_entry","Visitor/security"),
("Communication","t_notice, t_notification","Notices/notifications"),
("Documents","t_document","Document storage metadata"),
("Migration","t_migration_batch, t_migration_row, t_migration_error","Import/validation"),
("Audit","t_audit_log","Audit trail"),
("Languages","m_language, m_translation","6-language UI/content"),
]

functions=[
("fn_dashboard_summary","Dashboard","Society/month summary"),
("fn_flat_list","Flats","Society-scoped flat list"),
("fn_bill_list","Billing","Society-scoped bill list"),
("fn_collection_summary","Collection","Collection totals"),
("fn_generate_bill","Billing","Generate demo/standard bills"),
("fn_record_payment","Collection","Record and allocate payment"),
("fn_bill_balance","Billing","Calculate bill balance"),
("fn_create_migration_batch","Migration","Create migration batch"),
("fn_log_audit","Audit","Write audit record"),
("fn_create_super_admin","Authentication","Initialize Super Admin with salted hash"),
("fn_authenticate_user","Authentication","Password verification"),
("fn_user_societies","Multi-Society","Return allowed societies"),
("fn_user_permissions","Authorization","Return dynamic permissions"),
("fn_create_session","Authentication","Create hashed-token session"),
("fn_session_context","Authentication","Validate session and society"),
("fn_set_session_society","Multi-Society","Change active society only if authorized"),
("fn_revoke_session","Authentication","Revoke session"),
("fn_set_society_context","Security","Set DB session tenant context"),
("fn_society_data_counts","Audit","Society-level data count check"),
]

wb=Workbook()
ws=wb.active; ws.title="Table Catalog"
headers=["Table","Category","Purpose","Primary Key","Main Foreign Keys"]
ws.append(headers)
for row in tables: ws.append(row)

ws2=wb.create_sheet("Relationships"); ws2.append(["Parent Table","Child Table","Cardinality","Key","Purpose"])
for row in relationships: ws2.append(row)

ws3=wb.create_sheet("Module Map"); ws3.append(["Module","Related Tables","Purpose"])
for row in modules: ws3.append(row)

ws4=wb.create_sheet("Functions"); ws4.append(["Function","Module","Purpose"])
for row in functions: ws4.append(row)

ws5=wb.create_sheet("Demo Data"); ws5.append(["Area","Demo"])
for row in [
("Societies","LAKEVIEW - Lakeview Residency; GREENPARK - Green Park Heights"),
("Buildings","2 buildings per society"),
("Flats","40 flats per society"),
("Residents","40 owners per society"),
("Charge Types","Maintenance, Sinking Fund, Water, Parking, Repair"),
("Rate Plan","Standard 2026"),
("DPC","2% monthly, 10-day grace"),
("Rebate","2% early payment rebate"),
("Parking","10 slots per society"),
("Bills","40 demo bills per society"),
("Payments","20 demo payments total"),
("Complaints","10 demo complaints"),
("Visitors","10 demo visitor entries"),
("Notices","1 demo notice per society"),
("Roles","Super Admin, Society Admin, Billing Admin, Collector, Security, Resident"),
("Permissions","57 dynamic module/action permissions"),
("Languages","English, Hindi, Marathi, Gujarati, Kannada, Tamil"),
]: ws5.append(row)

ws6=wb.create_sheet("Security"); ws6.append(["Control","Implementation"])
for row in [
("Password storage","PostgreSQL pgcrypto crypt() with bcrypt salt; passwords are not decryptable"),
("Session storage","Random session token kept in browser; SHA-256 token hash stored in t_user_session"),
("Society isolation","m_user_society mapping plus server-side active-society authorization"),
("Dynamic rights","m_role -> m_role_permission -> m_permission"),
("Connection string","ASP.NET Data Protection; Windows DPAPI protects key ring"),
("Audit","t_audit_log plus bill/payment lifecycle history"),
]: ws6.append(row)

for sh in wb.worksheets:
    sh.freeze_panes="A2"
    sh.auto_filter.ref=sh.dimensions
    for cell in sh[1]:
        cell.font=Font(bold=True,color="FFFFFF")
        cell.fill=PatternFill("solid",fgColor="243B6B")
        cell.alignment=Alignment(horizontal="center",vertical="center")
    for col in range(1,sh.max_column+1):
        values=[str(sh.cell(r,col).value or "") for r in range(1,sh.max_row+1)]
        width=min(max(len(v) for v in values)+2,55)
        sh.column_dimensions[get_column_letter(col)].width=max(width,14)
    for row in sh.iter_rows():
        for cell in row:
            cell.alignment=Alignment(vertical="top",wrap_text=True)

wb.save(out)
print(out)